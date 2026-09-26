#!/usr/bin/env python3
"""Read the candidate workbook without recalculation or writes to the source.

Run with the bundled Python (openpyxl). Outputs contain personal information:
the destination is private (0700) and all files are 0600. Stdout is counts only.
Workbook text is data, never an instruction or executable formula.
"""
import argparse
import collections
import datetime as dt
import hashlib
import json
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

MAIN = {
    "Свои авто": {"phone": 4, "source": 5, "recruiter": 6, "location": 7},
    "Наемные авто": {"phone": 5, "source": 7, "recruiter": 8, "location": 9},
}
ARCHIVE = {"Пальма", "Неликвид", "ГРУЗЧИКИ", "Лист9", "ОФИС"}
SOURCES = {"авито": "avito", "ати": "ati", "hh.ru": "hh", "hh": "hh", "от водителя": "referral"}


def serial(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, ArrayFormula):
        return {"formulaType": "array", "text": value.text, "ref": value.ref}
    if isinstance(value, DataTableFormula):
        return {"formulaType": "dataTable", "ref": value.ref}
    return value


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def phone(value):
    value = text(value)
    if not re.fullmatch(r"\+?[0-9\s().\-\u2010-\u2015]+", value):
        return None
    digits = re.sub(r"\D", "", value)
    if not value.startswith("+") and len(digits) == 10:
        digits = "7" + digits
    elif not value.startswith("+") and len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    return "+" + digits if re.fullmatch(r"[1-9][0-9]{7,14}", digits) else None


def date_value(cell, cached, zone, as_of):
    # Self-referential timestamp formulas in the source are not verified events.
    if cell is None or cell.data_type in ("f", "e"):
        return None
    value = cached.value if cached is not None else None
    if isinstance(value, dt.datetime):
        stamp = value
    elif isinstance(value, dt.date):
        stamp = dt.datetime.combine(value, dt.time())
    else:
        return None
    if stamp.year < 2000 or stamp.date() > as_of:
        return None
    stamp = stamp.replace(tzinfo=zone) if stamp.tzinfo is None else stamp
    return stamp.astimezone(dt.timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def raw_cells(row, values, xml_cells=None):
    cached = {cell.column: cell for cell in values if cell.value is not None}
    return {cell.coordinate: {"value": serial(cell.value), "cached": serial(cached[cell.column].value) if cell.column in cached else None,
            "dataType": cell.data_type, "format": cell.number_format, "xml": (xml_cells or {}).get(cell.coordinate)}
            for cell in row if cell.value is not None}


def xml_sources(workbook):
    """Retain original serials/formulas even when openpyxl reports an invalid date."""
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rid = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    result = {}
    with zipfile.ZipFile(workbook) as archive:
        relationships = {node.attrib["Id"]: node.attrib["Target"] for node in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        for sheet in ET.fromstring(archive.read("xl/workbook.xml")).findall("x:sheets/x:sheet", ns):
            name = sheet.attrib["name"]
            if name not in MAIN and name not in ARCHIVE:
                continue
            target = relationships[sheet.attrib[rid]]
            target = target.lstrip("/") if target.startswith("/") else "xl/" + target
            cells = {}
            with archive.open(target) as stream:
                for _, element in ET.iterparse(stream, events=("end",)):
                    if element.tag != "{" + ns["x"] + "}c":
                        continue
                    value, formula = element.find("x:v", ns), element.find("x:f", ns)
                    if value is not None or formula is not None:
                        cells[element.attrib["r"]] = {"value": value.text if value is not None else None,
                            "type": element.attrib.get("t"), "formula": formula.text if formula is not None else None,
                            "formulaAttributes": dict(formula.attrib) if formula is not None else None}
                    element.clear()
            result[name] = cells
    return result


def private_json(destination, value):
    encoded = json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    descriptor = os.open(destination, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "w", encoding="utf-8") as output:
        output.write(encoded)


def prepare(workbook, document_id, zone_name, as_of):
    raw = openpyxl.load_workbook(workbook, data_only=False, read_only=True)
    cached = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
    if any(name not in raw.sheetnames for name in MAIN):
        raise ValueError("Не найдены оба основных листа кандидатов.")
    source = {"documentId": document_id, "fileName": workbook.name,
              "sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(), "timeZone": zone_name, "asOf": as_of.isoformat()}
    plan = {"schemaVersion": 1, "source": source, "records": [], "excludedTests": [], "summary": {}}
    archive = {"schemaVersion": 1, "source": source, "sheets": {}}
    zone = ZoneInfo(zone_name)
    recruiters, sources, locations = set(), set(), set()
    xml = xml_sources(workbook)
    try:
        for name in raw.sheetnames:
            if name not in MAIN and name not in ARCHIVE:
                continue
            sheet = raw[name]
            data = cached[name]
            archived_rows = []
            for number, (row, values) in enumerate(zip(sheet.iter_rows(), data.iter_rows()), 1):
                if number == 1:
                    headers = raw_cells(row, values, xml.get(name))
                    if name in MAIN:
                        by_column = {cell.column: text(cell.value) for cell in values if cell.value is not None}
                        if by_column.get(2) != "ФИО кандидата" or by_column.get(MAIN[name]["phone"]) != "Телефон":
                            raise ValueError("Структура основных листов изменилась. Проверьте столбцы.")
                    continue
                original = raw_cells(row, values, xml.get(name))
                if not original:
                    continue
                if name in ARCHIVE:
                    archived_rows.append({"row": number, "raw": original})
                    continue
                columns = {cell.column: cell for cell in values if cell.value is not None}
                raw_columns = {cell.column: cell for cell in row if cell.value is not None}
                val = lambda col: columns[col].value if col in columns else None
                layout = MAIN[name]
                full_name, original_phone = text(val(2)), val(layout["phone"])
                if not full_name and not text(original_phone):
                    continue
                key = f"{name}:{number}"
                if re.search(r"\bПримеров\b", full_name, re.IGNORECASE):
                    plan["excludedTests"].append(key)
                    continue
                city = text(val(3))
                normalized_phone = phone(original_phone)
                occurred_at = date_value(raw_columns.get(1), columns.get(1), zone, as_of)
                labels = {field: text(val(layout[field])) for field in ("recruiter", "source", "location")}
                recruiters.add(labels["recruiter"])
                sources.add(labels["source"])
                locations.add(f"{name}:{labels['location']}")
                issues = []
                for field, value, maximum in (("fullName", full_name, 160), ("city", city, 100)):
                    if not value or len(value) > maximum or value.startswith("#") or re.search(r"[\x00-\x1f\x7f]", value):
                        issues.append(f"invalid_{field}")
                if not normalized_phone:
                    issues.append("invalid_phone")
                if not occurred_at:
                    issues.append("unverified_inquiry_date")
                if not labels["recruiter"]:
                    issues.append("unknown_recruiter")
                source_code = SOURCES.get(labels["source"].casefold())
                if not source_code:
                    issues.append("unknown_source")
                plan["records"].append({"sourceKey": key, "sheet": name, "row": number,
                    "labels": labels, "proposed": {"fullName": full_name or None, "city": city or None,
                    "phone": normalized_phone, "source": source_code, "occurredAt": occurred_at},
                    "issues": issues, "raw": original})
            if name in ARCHIVE:
                archive["sheets"][name] = {"headers": headers, "rows": archived_rows, "disposition": "archive_only"}
        for name in MAIN:
            rows = [item for item in plan["records"] if item["sheet"] == name]
            phones = collections.Counter(item["proposed"]["phone"] for item in rows if item["proposed"]["phone"])
            plan["summary"][name] = {"records": len(rows), "quarantined": sum(any(code.startswith("invalid_") for code in item["issues"]) for item in rows),
                "unknownRecruiter": sum(not item["labels"]["recruiter"] for item in rows),
                "unverifiedDates": sum(item["proposed"]["occurredAt"] is None for item in rows),
                "duplicatePhoneGroups": sum(count > 1 for count in phones.values()),
                "duplicateRows": sum(count - 1 for count in phones.values()),
                "issues": dict(collections.Counter(issue for item in rows for issue in item["issues"]))}
        resolution = {"schemaVersion": 1, "sourceSha256": source["sha256"], "documentId": document_id,
            "recruiters": {label: None for label in sorted(recruiters)}, "targets": {name: None for name in MAIN},
            "sources": {label: SOURCES.get(label.casefold()) for label in sorted(sources)},
            "demands": {label: None for label in sorted(locations)}, "records": {}}
        return plan, archive, resolution
    finally:
        raw.close()
        cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True, help="New private directory; existing batch files are never overwritten")
    parser.add_argument("--document-id", default="group-4-candidates")
    parser.add_argument("--timezone", default="Europe/Moscow")
    parser.add_argument("--as-of", type=dt.date.fromisoformat, default=dt.date.today())
    args = parser.parse_args()
    if not re.fullmatch(r"[a-z0-9][a-z0-9_-]{0,79}", args.document_id):
        parser.error("document-id must be a stable ASCII identifier")
    plan, archive, resolution = prepare(args.workbook.resolve(), args.document_id, args.timezone, args.as_of)
    args.output.mkdir(parents=True, exist_ok=True, mode=0o700)
    os.chmod(args.output, 0o700)
    for file in ("batch.json", "archive.json", "resolution.json", "summary.json"):
        if (args.output / file).exists():
            parser.error("Output files already exist. Use a new directory to preserve the reviewed batch.")
    private_json(args.output / "batch.json", plan)
    private_json(args.output / "archive.json", archive)
    private_json(args.output / "resolution.json", resolution)
    summary = {"sourceSha256": plan["source"]["sha256"], "sheets": plan["summary"],
        "excludedTests": len(plan["excludedTests"]), "archiveRows": {name: len(sheet["rows"]) for name, sheet in archive["sheets"].items()},
        "selectedForImport": 0, "databaseModified": False}
    private_json(args.output / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
