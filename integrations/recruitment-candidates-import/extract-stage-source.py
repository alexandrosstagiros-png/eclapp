#!/usr/bin/env python3
"""Extract source headers, other sheets and excluded examples without evaluating formulas.

The workbook is read-only. Output contains private source data (0600).
"""
import argparse
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("candidate_prepare", HERE / "prepare.py")
prepare = importlib.util.module_from_spec(spec)
spec.loader.exec_module(prepare)


def xml_cells(book):
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    relationship = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    result = {}
    with zipfile.ZipFile(book) as archive:
        relationships = {node.attrib["Id"]: node.attrib["Target"] for node in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        for sheet in ET.fromstring(archive.read("xl/workbook.xml")).findall("x:sheets/x:sheet", ns):
            target = relationships[sheet.attrib[relationship]]
            target = target.lstrip("/") if target.startswith("/") else "xl/" + target
            cells = {}
            with archive.open(target) as stream:
                for _, element in ET.iterparse(stream, events=("end",)):
                    if element.tag != "{" + ns["x"] + "}c":
                        continue
                    value, formula = element.find("x:v", ns), element.find("x:f", ns)
                    if value is not None or formula is not None:
                        cells[element.attrib["r"]] = {"value": value.text if value is not None else None,
                            "type": element.attrib.get("t"), "formula": formula.text if formula is not None else None,
                            "formulaAttributes": dict(formula.attrib) if formula is not None else None}
                    element.clear()
            result[sheet.attrib["name"]] = cells
    return result


def extract(workbook, batch):
    source_hash = hashlib.sha256(workbook.read_bytes()).hexdigest()
    if source_hash != batch["source"]["sha256"] or workbook.name != batch["source"]["fileName"]:
        raise ValueError("Source snapshot mismatch")
    raw = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    xml = xml_cells(workbook)
    result = {"schemaVersion": 1, "source": batch["source"], "headers": {}, "sheets": {}, "excludedTests": [], "mainRemainders": []}
    excluded = set(batch.get("excludedTests", []))
    main_keys = {record["sourceKey"] for record in batch["records"]}
    try:
        for name in raw.sheetnames:
            extra = name not in prepare.MAIN and name not in prepare.ARCHIVE
            rows = []
            for number, (row, values) in enumerate(zip(raw[name].iter_rows(), cached[name].iter_rows()), 1):
                original = prepare.raw_cells(row, values, xml.get(name))
                if number == 1:
                    result["headers"][name] = original
                if not original:
                    continue
                if extra:
                    # Include row 1: a headers-only report sheet is still source data.
                    rows.append({"row": number, "raw": original})
                elif f"{name}:{number}" in excluded:
                    result["excludedTests"].append({"sheet": name, "row": number, "raw": original})
                elif name in prepare.MAIN and number > 1 and f"{name}:{number}" not in main_keys:
                    # Rows without an identity may still contain comments, dates,
                    # false flags or formulas. Preserve them as source-only archive.
                    result["mainRemainders"].append({"sheet": name, "row": number, "raw": original})
                if sum(len(sheet["rows"]) for sheet in result["sheets"].values()) + len(rows) + len(result["mainRemainders"]) > 50000:
                    raise ValueError("Too many source rows")
            if extra:
                result["sheets"][name] = {"headers": result["headers"][name], "rows": rows, "disposition": "archive_only"}
        if {f"{row['sheet']}:{row['row']}" for row in result["excludedTests"]} != excluded:
            raise ValueError("Excluded source rows are missing")
        if hashlib.sha256(workbook.read_bytes()).hexdigest() != source_hash:
            raise ValueError("Source changed while reading")
        return result
    finally:
        raw.close()
        cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--directory", type=Path, required=True)
    args = parser.parse_args()
    batch = json.loads((args.directory / "batch.json").read_text())
    result = extract(args.workbook.resolve(), batch)
    remaining = {"schemaVersion": 1, "source": result["source"], "records": result.pop("mainRemainders")}
    outputs = [(args.directory / "extra-archive.json", result), (args.directory / "main-remainders.json", remaining)]
    for output, value in outputs:
        if output.exists() and json.loads(output.read_text()) != value:
            raise ValueError("Existing extraction differs; keep the reviewed snapshot")
    for output, value in outputs:
        if not output.exists():
            with os.fdopen(os.open(output, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600), "w") as file:
                json.dump(value, file, ensure_ascii=False, indent=2)
                file.write("\n")
    print(json.dumps({"sourceSha256": result["source"]["sha256"], "extraSheets": len(result["sheets"]),
        "extraRows": sum(len(sheet["rows"]) for sheet in result["sheets"].values()), "excludedExamples": len(result["excludedTests"]),
        "mainRemainders": len(remaining["records"])}))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        raise SystemExit("Source extraction failed; workbook and existing files preserved.")
