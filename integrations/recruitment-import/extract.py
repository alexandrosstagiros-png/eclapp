#!/usr/bin/env python3
"""Read both recruitment sheets, retaining raw values, cached values and merges.

Requires the bundled Python runtime with openpyxl. Never modifies the workbook.
The output contains internal contacts and must remain private.
"""
import argparse
import datetime
import hashlib
import json
import os
from pathlib import Path

import openpyxl


def serializable(value):
    if isinstance(value,(datetime.datetime,datetime.date,datetime.time)):
        return value.isoformat()
    return value


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args = parser.parse_args()
    workbook = args.workbook.resolve()
    raw = openpyxl.load_workbook(workbook,data_only=False)
    cached = openpyxl.load_workbook(workbook,data_only=True)
    result = {'sourceFile':str(workbook),'sha256':hashlib.sha256(workbook.read_bytes()).hexdigest(),
              'sheets':{}}
    for name in ('Проекты','Перевозчики'):
        sheet = raw[name]
        rows = []
        for row in sheet:
            cells = {}
            for cell in row:
                if cell.value is None:
                    continue
                cells[cell.coordinate] = {'value':serializable(cell.value),
                    'cached':serializable(cached[name][cell.coordinate].value),
                    'format':cell.number_format,'dataType':cell.data_type,
                    'hyperlink':cell.hyperlink.target if cell.hyperlink else None}
            if cells:
                rows.append({'row':row[0].row,'cells':cells})
        result['sheets'][name] = {'rows':rows,
            'mergedRanges':sorted(str(m) for m in sheet.merged_cells.ranges),
            'hiddenRows':[r for r,d in sheet.row_dimensions.items() if d.hidden],
            'hiddenColumns':[c for c,d in sheet.column_dimensions.items() if d.hidden],
            'maxRow':sheet.max_row,'maxColumn':sheet.max_column}
    raw.close()
    cached.close()
    args.output.parent.mkdir(parents=True,exist_ok=True)
    descriptor = os.open(args.output,os.O_WRONLY|os.O_CREAT|os.O_TRUNC,0o600)
    try:
        os.fchmod(descriptor,0o600)
        with os.fdopen(descriptor,'w') as output:
            output.write(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    except BaseException:
        try:
            os.close(descriptor)
        except OSError:
            pass
        raise
    print(json.dumps({'sha256':result['sha256'],
                      'sheets':{name:{'rows':len(sheet['rows']),'merges':len(sheet['mergedRanges'])}
                                for name,sheet in result['sheets'].items()}},ensure_ascii=False))


if __name__ == '__main__':
    main()
